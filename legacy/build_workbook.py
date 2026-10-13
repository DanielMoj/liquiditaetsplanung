from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

YEL = PatternFill("solid", fgColor="FFFF00")   # Eingabe: du setzt den Wert
BLU = PatternFill("solid", fgColor="DDEBF7")   # Pflichtdaten: Ist/BWA + Klassifikation (beobachtet/strukturell)
GRY = PatternFill("solid", fgColor="D9D9D9")   # von der Engine erzeugt / strukturell fix (zaehlt in die Rechnung)
CTX = PatternFill("solid", fgColor="EDE7DD")   # reiner Kontext: von der Engine NICHT gelesen
AR  = Font(name="Arial", size=10)
ARB = Font(name="Arial", size=10, bold=True)
ARI = Font(name="Arial", size=10, italic=True)   # Kontext: kursiv -> "inaktiv/Info"

# Anzeigeformate (rein visuell; gespeicherte Werte bleiben unveraendert -> R liest weiter korrekt).
# Auf deutschem Excel rendert #,##0 mit Tausenderpunkt und Dezimalkomma.
EUR  = '#,##0" €"'      # Geldbetrag, 0 Nachkommastellen, Tausenderpunkt
PCT  = '0.0%'           # Rate/Anteil/Wahrscheinlichkeit (Wert 0,038 -> 3,8%)
CORR = '0.00'           # Korrelation, dimensionslos (0,40 / 1,00)
DAYS = '0" Tage"'       # Zahlungsziel / Verzug
MON  = '0" Mon."'       # Monate (Laufzeit, hmax, USt-Lag)
HEB  = '0"%"'           # Hebesatz: Wert ist schon in Prozentpunkten (400 = 400%), nicht *100
DATE = 'DD.MM.YYYY'     # Faelligkeitsdatum

def put(ws, r, c, v, yellow=False, blue=False, grey=False, context=False, bold=False, text=False, fmt=None):
    cell = ws.cell(r, c)
    cell.value = v
    cell.font = ARB if bold else (ARI if context else AR)
    if yellow:    cell.fill = YEL
    elif blue:    cell.fill = BLU
    elif grey:    cell.fill = GRY
    elif context: cell.fill = CTX
    if text and v is not None:
        cell.number_format = "@"
    elif fmt is not None:
        cell.number_format = fmt
    return cell

def header_sheet(ws, title, subtitle, header):
    put(ws, 1, 1, title, bold=True)
    put(ws, 2, 1, subtitle)
    # row 3 blank -> header lands on row 4 (engine uses read_excel(skip=3))
    for j, h in enumerate(header, start=1):
        put(ws, 4, j, h, bold=True, text=True)

wb = Workbook()

# ---------- 00_Anleitung ----------
ws = wb.active; ws.title = "00_Anleitung"
lines = [
 ("Liquiditätsmodell — Mandats-Eingabemappe", True),
 ("Eine Mappe je Mandat. Blätter 01–08 befüllen, dann die R-Skripte laufen lassen.", False),
 ("", False),
 ("Ablauf", True),
 ("1. 01_Positionen: BWA-/GuV-Istwerte je Monat + Klassifikation (Gruppe, Zahlungsprofil, Treiber, USt-Satz).", False),
 ("2. 02_Zahlungsprofile: Forderungs-/Verbindlichkeitslaufzeiten — aus OPOS-Listen schätzen, nicht raten.", False),
 ("3. 03_Stammdaten & Eröffnung: Eröffnungssalden zum Forecast-Beginn (Zahlungsmittel, KK-Linie …).", False),
 ("4. 04_Annahmen: vorwärtsgerichtete Treiber (Menge, Preis, Tarif …) — die gelben Felder.", False),
 ("5. 05_Finanzierungsplan: Investitionen, Fördermittel, Tilgung, Darlehensaufnahme je Forecast-Monat.", False),
 ("6. 08_Steuern: Rechtsform + Steuerstatus + Sätze — entscheidet, welche Steuern zahlungswirksam sind.", False),
 ("7. 09_Darlehen: einzelne Kredite mit Tilgungsplan (optional) — Zins wirkt in GuV und Liquidität.", False),
 ("8. 10_OPOS: offene Forderungen/Verbindlichkeiten mit Zahlungsziel (optional) — ersetzt im Nahbereich die statistischen Profile.", False),
 ("", False),
 ("Prinzip", True),
 ("• Daten (beobachtet) und Annahmen (Urteil) sind getrennt — siehe Farb-Legende unten. Gelb = Eingabe, Blau = Pflichtdaten, Grau = von der Engine erzeugt.", False),
 ("• Die Engine validiert zuerst: rechnet sie die Ist-Historie nach? Erst dann der Forecast.", False),
 ("• Engine + Klassifikation bleiben gleich; für ein anderes Mandat nur die Daten und Annahmen tauschen.", False),
]
for i,(t,b) in enumerate(lines, start=1):
    put(ws, i, 1, t, bold=b)
ws.column_dimensions["A"].width = 100

# ---- Farb-Legende (Swatch = die echte Fuellfarbe) ----
base = len(lines) + 2
put(ws, base,   1, "Farbcodierung der Felder", bold=True)
put(ws, base+1, 1, "Gelb — Eingabe: du setzt den Wert. Annahmen (04), Finanzierungsplan (05), Eröffnungssalden/KK-Linie/Steuersätze (03/08), Bandbreiten p_min/p_max (06), Darlehen (09).", yellow=True)
put(ws, base+2, 1, "Blau — Pflichtdaten, beobachtet/strukturell: die Ist-/BWA-Monatswerte und die Klassifikation (Gruppe, Zahlungsprofil, Treiber, USt) in 01_Positionen.", blue=True)
put(ws, base+3, 1, "Grau — von der Engine erzeugt oder strukturell fix; du lässt es stehen, es zählt aber in die Rechnung: geschätzte cv (menge/material in 06), Diagonale 1,0 in 07.", grey=True)
put(ws, base+4, 1, "Beige/kursiv — reiner Kontext: von der Engine NICHT gelesen. Nur Orientierung; ändern oder leeren ändert das Ergebnis nicht (untere Zeilen in 03).", context=True)
put(ws, base+5, 1, "Ohne Farbe — Beschriftungen, Schlüssel, Hinweise, Kopfzeilen. Nicht ausfüllen.")
put(ws, base+7, 1, "Hinweis: Farben sind rein visuell. Die R-Engine liest die Werte unverändert; read_excel ignoriert Zellfarben.")

# ---------- 01_Positionen (+ ust) ----------
ws = wb.create_sheet("01_Positionen")
months = [f"2025-{m:02d}" for m in range(1,13)]
header_sheet(ws, "01_Positionen — Ist-Werte je Monat + Klassifikation",
             "Klassifikation (Spalten C–F) = Mapping. Monatsspalten = gebuchte Ist-Werte (EUR). Aufwand positiv. ust = USt-Satz je Zeile (0 / 0.07 / 0.19), nur bei ust_status=pflichtig genutzt.",
             ["pos","label","grp","cashprofile","driver","ust"] + months)
def row12(v): return [v]*12
pos_rows = [
 ("r_drg","DRG-Erlöse (Kassen)","umsatz","kassen","menge_preis",[8111200,7718100,8111900,7502200,7546500,7223300,6692000,6281700,7354900,7647100,7794600,7012700]),
 ("r_ze","Zusatzentgelte","umsatz","kassen","menge_preis",[865900,862100,927700,853100,814700,805700,751500,688400,845200,888900,872000,761700]),
 ("r_pfb","Pflegebudget","umsatz","abschlag","pflegebudget",[3828500,3721700,3766700,3658500,3704900,3641300,3552400,3507300,3670700,3793100,3726400,3525500]),
 ("r_wahl","Wahlleistungen","umsatz","wahl","menge_preis",[511300,491600,522400,471100,467700,420000,392400,385900,434700,468400,466400,421300]),
 ("r_amb","Ambulante Erlöse","umsatz","amb","menge_preis",[1680800,1618700,1680400,1540500,1516500,1525200,1370700,1336000,1491100,1644700,1595700,1436900]),
 ("r_best","Bestandsveränderung","bestand","noncash","menge",[180000,-90000,60000,120000,-40000,-160000,-210000,30000,140000,90000,110000,-80000]),
 ("m_arz","Arzneimittel","material","dpo","material",[1343300,1291600,1333700,1246200,1233200,1179200,1046000,1010000,1183000,1251100,1258600,1166000]),
 ("m_sach","Med. Sachbedarf","material","dpo","material",[2070700,1965400,2081700,1900500,1921600,1795600,1626700,1588000,1863100,1940700,1997600,1804600]),
 ("m_leb","Lebensmittel","material","dpo","menge",[239400,235300,247500,235800,229500,226800,220500,222800,236100,237300,233800,222900]),
 ("o_son","Sonstige betr. Erlöse","betr_ertrag","inmonth","infl",[573900,543600,572300,565100,509900,556000,553500,544100,526500,550800,516400,537400]),
 ("o_sopo","Auflösung Sonderposten","sopo","noncash","flat",row12(500000)),
 ("p_arz","Ärztlicher Dienst","personal","inmonth","tarif",[2203000,2225000,2221600,2194600,2205000,2190800,2223200,2221300,2196100,2211100,3753200,2189700]),
 ("p_pfl","Pflegedienst","personal","inmonth","tarif",[3483200,3467700,3484300,3467100,3430400,3452800,3431700,3494700,3491200,3488000,5862300,3434400]),
 ("p_mtd","Med.-techn. Dienst","personal","inmonth","tarif",[872800,874000,859100,866000,858800,870800,870900,859800,865800,867100,1464300,872700]),
 ("p_fun","Funktionsdienst","personal","inmonth","tarif",[707700,704700,709300,712000,704500,706100,715800,710900,707900,709000,1194600,704800]),
 ("p_wir","Wirtschaftsdienst","personal","inmonth","tarif",[235500,236700,235000,234900,234200,236900,235000,238200,237900,234200,399100,237000]),
 ("p_vw","Verwaltungsdienst","personal","inmonth","tarif",[352300,351800,357500,354900,354200,356400,356600,352200,351500,353900,601000,354100]),
 ("p_tec","Technischer Dienst","personal","inmonth","tarif",[134500,134300,135200,132800,133600,133400,134800,133200,133000,133700,227000,133300]),
 ("p_son","Sonderdienste","personal","inmonth","tarif",[140300,144100,141400,143800,142000,139200,144600,143700,144400,144200,244100,139900]),
 ("k_ene","Wasser/Energie/Brennstoffe","sachkosten","abschlag","energie",[784300,736600,643600,491900,401300,340800,290900,323100,404400,537500,708000,864100]),
 ("k_rei","Reinigung","sachkosten","inmonth","infl",[301000,303900,293800,296300,308400,301400,300800,304500,292000,301500,300100,306300]),
 ("k_mie","Mieten/Leasing","sachkosten","inmonth","flat",row12(200000)),
 ("k_ver","Versicherungen/Beiträge","sachkosten","annual_jan","flat",row12(208300)),
 ("k_ins","Instandhaltung","sachkosten","inmonth","infl",[290900,274400,370300,542100,613600,663700,545400,454900,408300,336100,350600,307700]),
 ("k_it","IT/EDV","sachkosten","inmonth","infl",[281500,285200,283800,293200,276600,288200,277400,281000,287200,278800,279200,289000]),
 ("k_bue","Verwaltungs-/Bürobedarf","sachkosten","inmonth","infl",[111700,116200,122500,122500,111700,113300,113900,121700,121100,121100,115100,112700]),
 ("k_fre","Fremdleistungen","sachkosten","inmonth","menge",[539400,522800,537600,517000,490800,441200,438800,407300,486300,529800,493600,436900]),
 ("k_son","Sonstige betr. Aufwendungen","sachkosten","inmonth","infl",[254100,268400,259400,270000,273600,257200,271000,259100,266300,279600,277700,253600]),
 ("k_afa","Abschreibungen","afa","noncash","flat",row12(666700)),
 ("n_zie","Zinserträge","zins_e","inmonth","flat",[8200,8000,7500,8800,8600,7900,8700,8400,8200,7500,7600,9000]),
 ("n_zia","Zinsaufwand","zins_a","quarterly","flat",[107500,105300,107100,111400,114200,119800,132600,143400,130600,119200,113500,103200]),
 ("n_nee","Neutrale Erträge","neutral_e","noncash","flat",[48700,48000,68500,86500,99100,421700,64400,69500,86400,104800,137500,813200]),
 ("n_nea","Neutrale Aufwendungen","neutral_a","noncash","flat",[35700,44600,39600,47100,45400,46200,45600,38800,45300,36800,46300,46100]),
 ("t_st","Sonstige Steuern (Grund-/Kfz-)","steuer","inmonth","flat",[11100,14100,12300,11500,14000,11100,10100,11000,11600,14300,14800,11400]),
]
r = 5
for pos,label,grp,cp,drv,vals in pos_rows:
    put(ws,r,1,pos); put(ws,r,2,label)
    put(ws,r,3,grp,blue=True); put(ws,r,4,cp,blue=True); put(ws,r,5,drv,blue=True)   # Klassifikation = strukturelle Pflichtangabe
    put(ws,r,6,0,blue=True,fmt=PCT)  # ust = 0 (Krankenhaus: steuerfrei nach §4 Nr.14)
    for j,v in enumerate(vals, start=7): put(ws,r,j,v,blue=True,fmt=EUR)  # Ist-/BWA-Werte
    r += 1
for col,w in {"A":8,"B":28,"C":12,"D":12,"E":13,"F":7}.items():
    ws.column_dimensions[col].width = w
for j in range(7,19):
    ws.column_dimensions[get_column_letter(j)].width = 13

# ---------- 02_Zahlungsprofile ----------
ws = wb.create_sheet("02_Zahlungsprofile")
header_sheet(ws,"02_Zahlungsprofile — Anteil der Zahlung je Monat ab Leistung/Rechnung",
             "Aus OPOS-Aging ableiten. f1 = Anteil im Monat, f2 = Folgemonat usw. (gelb = editierbar).",
             ["profile","f1","f2","f3","f4","Anwendung"])
prof = [
 ("kassen",0.6,0.35,0.05,None,"DRG / Zusatzentgelte (gesetzl. Kassen)"),
 ("wahl",0.3,0.4,0.25,0.05,"Wahlleistungen / Selbstzahler"),
 ("amb",0.1,0.2,0.4,0.3,"Ambulant (KV-Quartalsabrechnung)"),
 ("dpo",0.4,0.6,None,None,"Material / Lieferanten"),
]
r=5
for name,f1,f2,f3,f4,anw in prof:
    put(ws,r,1,name)
    for j,v in enumerate([f1,f2,f3,f4], start=2):
        if v is not None: put(ws,r,j,v,yellow=True,fmt=PCT)
    put(ws,r,6,anw)
    r+=1
put(ws,r+1,1,"Schlüsselwörter (kein Lag): inmonth · noncash · annual_jan · abschlag · quarterly")
ws.column_dimensions["A"].width=12; ws.column_dimensions["F"].width=40

# ---------- 03_Stammdaten_Eroeffnung ----------
ws = wb.create_sheet("03_Stammdaten_Eroeffnung")
header_sheet(ws,"03_Stammdaten & Eröffnung",
             "Gelb = von der Engine gelesen. Eröffnungssalden zum Forecast-Beginn.",
             ["key","wert","Hinweis"])
stamm = [
 ("Mandant","Krankenhaus gGmbH (Beispiel)","",False,False),
 ("Branche","Krankenhaus · Grund-/Regelversorgung","",False,False),
 ("Ist-Zeitraum","2025-01 bis 2025-12","",False,True),
 ("Forecast-Beginn","2026-01","",True,True),
 ("open_cash",-2586600,"Zahlungsmittel 01.01. (= Endbestand Vorjahr)",True,False),
 ("kk_limit",12000000,"Kontokorrentlinie",True,False),
 ("kk_zins",0.065,"KK-Zinssatz p.a.",True,False),
 ("mindestreserve",1000000,"angestrebte Mindestliquidität",True,False),
 ("Forderungen",20000000,"Kontext (nicht gerechnet)",False,False),
 ("Verbindlichkeiten",8000000,"Kontext",False,False),
 ("Vorraete",3500000,"Kontext",False,False),
 ("Rueckstellungen",14000000,"Kontext",False,False),
 ("Darlehen",28000000,"Kontext",False,False),
]
r=5
_fmt03 = {"open_cash":EUR,"kk_limit":EUR,"mindestreserve":EUR,"Forderungen":EUR,
          "Verbindlichkeiten":EUR,"Vorraete":EUR,"Rueckstellungen":EUR,"Darlehen":EUR,"kk_zins":PCT}
for k,v,h,yel,txt in stamm:
    ctx = ("Kontext" in h)   # untere Zeilen: reine Info, von der Engine nicht gelesen
    put(ws,r,1,k,context=ctx)
    put(ws,r,2,v,yellow=yel,context=ctx,text=txt,fmt=_fmt03.get(k))
    put(ws,r,3,h,context=ctx)
    r+=1
ws.column_dimensions["A"].width=18; ws.column_dimensions["B"].width=28; ws.column_dimensions["C"].width=42

# ---------- 04_Annahmen ----------
ws = wb.create_sheet("04_Annahmen")
header_sheet(ws,"04_Annahmen — vorwärtsgerichtete Treiber (Szenario)",
             "Die Stellschrauben. In der Simulation (Phase 3+) wird jeder Wert durch eine Verteilung ersetzt.",
             ["key","wert","Bedeutung"])
ann = [
 ("menge",0.005,"Fallzahl-/Mengenentwicklung vs. Vorjahr"),
 ("preis",0.04,"Erlöspreis / Landesbasisfallwert"),
 ("tarif",0.038,"Tarifabschluss Personal"),
 ("energie",-0.05,"Energiepreis-Pfad"),
 ("material",0.03,"Materialkosteninflation (zus. zur Menge)"),
 ("pflegebudget",0.04,"Anpassung Pflegebudget"),
 ("inflation",0.03,"allg. Kostentrend (übrige Sachkosten)"),
 ("hmax",6,"Forecast-Monate"),
]
r=5
for k,v,b in ann:
    put(ws,r,1,k); put(ws,r,2,v,yellow=True,fmt=(MON if k=="hmax" else PCT)); put(ws,r,3,b)
    r+=1
ws.column_dimensions["A"].width=14; ws.column_dimensions["B"].width=10; ws.column_dimensions["C"].width=42

# ---------- 05_Finanzierungsplan ----------
ws = wb.create_sheet("05_Finanzierungsplan")
header_sheet(ws,"05_Finanzierungsplan — je Forecast-Monat (EUR)",
             "Auszahlungen/Zuflüsse außerhalb der GuV. Gelb = editierbar. M1 = erster Forecast-Monat.",
             ["Position"] + [f"M{i}" for i in range(1,13)])
fin = [
 ("Investitionen",[200000,200000,1500000,200000,200000,2000000,0,0,0,0,0,0]),
 ("Foerdermittel",[0,0,0,1500000,0,0,0,0,0,0,0,0]),
 ("Tilgung",row12(180000)),
 ("Darlehensaufnahme",row12(0)),   # 4 Mio. nach 09_Darlehen migriert (dort inkl. Zins/Tilgung)
]
r=5
for name,vals in fin:
    put(ws,r,1,name)
    for j,v in enumerate(vals, start=2): put(ws,r,j,v,yellow=True,fmt=EUR)
    r+=1
ws.column_dimensions["A"].width=20
for j in range(2,14): ws.column_dimensions[get_column_letter(j)].width=13

# ---------- 06_Verteilungen ----------
ws = wb.create_sheet("06_Verteilungen")
header_sheet(ws,"06_Verteilungen — Unsicherheit der ausgewählten Treiber (Phase 3)",
             "mode = Punktwert aus 04_Annahmen. p_min/p_max = epistemische Bandbreite (PERT). cv = aleatorische Monatsstreuung.",
             ["driver","p_min","p_max","cv","Hinweis"])
dist = [
 ("tarif",0.03,0.048,0,True,"Jahresrate; aleatorisch ~0 (jaehrlicher Schritt) — cv 0 von Natur aus"),
 ("preis",0.025,0.055,0,True,"Landesbasisfallwert; rein epistemisch — cv 0 von Natur aus"),
 ("pflegebudget",0.025,0.055,0,True,"Pflegebudget; rein epistemisch — cv 0 von Natur aus"),
 ("menge",-0.015,0.025,0.03,False,"cv wird in Phase 3 aus der Ist-Historie GESCHAETZT; dieser Wert nur Rueckfall bei <8 Monaten"),
 ("material",0,0.06,0.03,False,"cv wird in Phase 3 aus der Ist-Historie GESCHAETZT; dieser Wert nur Rueckfall bei <8 Monaten"),
]
r=5
for d,pmin,pmax,cv,cv_input,h in dist:
    put(ws,r,1,d); put(ws,r,2,pmin,yellow=True,fmt=PCT); put(ws,r,3,pmax,yellow=True,fmt=PCT)
    put(ws,r,4,cv,yellow=cv_input,grey=not cv_input,fmt=PCT); put(ws,r,5,h)   # cv: Eingabe(gelb) bzw. von Engine geschaetzt(grau)
    r+=1
put(ws,r+1,1,"Treiber ohne Zeile (energie, inflation) bleiben deterministisch.")
ws.column_dimensions["A"].width=14; ws.column_dimensions["E"].width=55

# ---------- 07_Korrelation ----------
ws = wb.create_sheet("07_Korrelation")
header_sheet(ws,"07_Korrelation — Abhängigkeit der epistemischen Treiber (Phase 4)",
             "Korrelation der Jahresannahmen (Gauß-Copula). Diagonale = 1, symmetrisch. Gelb = editierbar.",
             ["driver","tarif","preis","pflegebudget","menge","material"])
corr = [
 ("tarif",1,0.4,0.6,0,0.3),
 ("preis",0.4,1,0.5,0.1,0.2),
 ("pflegebudget",0.6,0.5,1,0,0.2),
 ("menge",0,0.1,0,1,0),
 ("material",0.3,0.2,0.2,0,1),
]
r=5
for row in corr:
    put(ws,r,1,row[0])
    for j,v in enumerate(row[1:], start=2):
        diag = (j-2 == r-5)
        put(ws,r,j,v,yellow=not diag,grey=diag,fmt=CORR)  # off-diagonal = Eingabe(gelb); Diagonale 1,0 = strukturell fix(grau)
    r+=1
ws.column_dimensions["A"].width=14

# ---------- 08_Steuern (NEU) ----------
ws = wb.create_sheet("08_Steuern")
header_sheet(ws,"08_Steuern — Rechtsform, Steuerstatus und Sätze (mandantenabhängig)",
             "Drei Schalter entscheiden, welche Steuer zahlungswirksam ist. Gelb = editierbar. Sätze als Parameter (Hebesatz je Gemeinde, KSt-Satz per Gesetz).",
             ["key","wert","Hinweis"])
tax = [
 ("rechtsform","kapital","kapital | personen | einzel  (transparent = personen/einzel)"),
 ("gewst_status","gemeinnuetzig","gewerblich | freiberuflich | gemeinnuetzig"),
 ("ust_status","frei","pflichtig | frei (§4 Nr.14 Heilbehandlung = frei)"),
 ("ust_methode","soll","soll (Versteuerung nach Leistung) | ist (nach Vereinnahmung, §20 UStG)"),
 ("kst_satz",0.15,"Körperschaftsteuersatz (nur Kapitalgesellschaft)"),
 ("soli_satz",0.055,"Solidaritätszuschlag auf die KSt"),
 ("gewst_messzahl",0.035,"Steuermesszahl Gewerbesteuer"),
 ("gewst_hebesatz",400,"Hebesatz in % (z.B. 400 = 400%)"),
 ("gewst_freibetrag",24500,"GewSt-Freibetrag (nur EK/Personengesellschaft)"),
 ("ust_lag",1,"Monate bis USt-Abführung (1; mit Dauerfristverlängerung 2)"),
 ("entnahme_quote",0.30,"Steuerentnahme als Anteil des Plangewinns (nur transparente Formen)"),
]
r=5
_fmt08 = {"kst_satz":PCT,"soli_satz":PCT,"gewst_messzahl":PCT,"entnahme_quote":PCT,
          "gewst_hebesatz":HEB,"gewst_freibetrag":EUR,"ust_lag":MON}
for k,v,h in tax:
    yel = k in ("rechtsform","gewst_status","ust_status","ust_methode","kst_satz","soli_satz","gewst_messzahl","gewst_hebesatz","gewst_freibetrag","ust_lag","entnahme_quote")
    txt = k in ("rechtsform","gewst_status","ust_status","ust_methode")
    put(ws,r,1,k); put(ws,r,2,v,yellow=yel,text=txt,fmt=_fmt08.get(k)); put(ws,r,3,h)
    r+=1
ws.column_dimensions["A"].width=18; ws.column_dimensions["B"].width=16; ws.column_dimensions["C"].width=60

# ---------- 09_Darlehen (NEU) ----------
ws = wb.create_sheet("09_Darlehen")
header_sheet(ws,"09_Darlehen — explizite Kredite mit Tilgungsplan (mandantenabhängig, optional)",
             "Je Zeile ein Kredit. Zins läuft in GuV (EBT→Ertragsteuer) und Liquidität, Tilgung/Auszahlung nur Liquidität. Gelb = editierbar. Blatt leer = altes Verhalten. tilgungsart: annuitaet | linear | endfaellig | nur_zins. start_m = Auszahlungsmonat (1..hmax); <=0 = bereits valutiert (Restschuld=betrag, ab M1 bedient). disagio mindert nur die Auszahlung.",
             ["id","betrag","zinssatz","tilgungsart","laufzeit_m","start_m","disagio"])
darlehen = [
 ("KfW_Invest", 4000000, 0.045, "annuitaet", 60, 6, 0),   # im 6. Forecast-Monat valutiert (= bisherige Darlehensaufnahme aus Blatt 05)
]
r=5
for id_,betrag,zins,art,lauf,start,dis in darlehen:
    put(ws,r,1,id_)
    put(ws,r,2,betrag,yellow=True,fmt=EUR); put(ws,r,3,zins,yellow=True,fmt=PCT); put(ws,r,4,art,yellow=True,text=True)
    put(ws,r,5,lauf,yellow=True,fmt=MON); put(ws,r,6,start,yellow=True,fmt=MON); put(ws,r,7,dis,yellow=True,fmt=PCT)
    r+=1
ws.column_dimensions["A"].width=16
for col,w in {"B":12,"C":10,"D":13,"E":12,"F":9,"G":9}.items(): ws.column_dimensions[col].width=w

# ---------- 10_OPOS (NEU) ----------
ws = wb.create_sheet("10_OPOS")
header_sheet(ws,"10_OPOS — offene Forderungen und Verbindlichkeiten mit Zahlungsziel (optional)",
             "Aus der OPOS-Liste. Je Zeile eine Einzelrechnung ODER ein Aging-Bucket (beides erlaubt). art = forderung | verbindlichkeit. betrag brutto. Termin per faellig_am (ISO-Datum) ODER zahlungsziel_tage ab Forecast-Beginn. wahrscheinlichkeit = Eingangs-/Erfüllungsquote (Haircut, default 1). verzug_tage = erwarteter Verzug bzw. Lieferantenstundung (default 0). ust_satz = eingebetteter USt-Satz (0 / 0.07 / 0.19), nur bei ust_methode=ist genutzt. Leeres Blatt = statistische Profile (Carry-in) wie bisher. Aktive Posten ersetzen den Carry-in der jeweiligen Seite (keine Doppelerfassung).",
             ["art","bezeichnung","betrag","faellig_am","zahlungsziel_tage","wahrscheinlichkeit","verzug_tage","ust_satz"])
# Datenzeilen bleiben leer (Beispiel-Mandant Krankenhaus: aus OPOS-Liste zu befuellen).
# Beispielzeile (zur Orientierung, auskommentiert):
#   forderung | Kasse Dezember | 5000000 | 2026-01-20 |        | 1.0  | 0  | 0
#   verbindlichkeit | Material 30-60 | 1500000 |       | 45     | 1.0  | 0  | 0.19
ws.column_dimensions["A"].width=15
# Kuenftige Eingabezeilen vorformatieren, damit eingetragene Werte gleich Einheiten zeigen
_fmt10 = {3:EUR, 4:DATE, 5:DAYS, 6:PCT, 7:DAYS, 8:PCT}
for rr in range(5,31):
    for cc,fm in _fmt10.items(): ws.cell(rr,cc).number_format = fm
for col,w in {"B":26,"C":12,"D":13,"E":17,"F":17,"G":12,"H":10}.items(): ws.column_dimensions[col].width=w

wb.save("Mandant_Liquiditaetsmodell_Vorlage.xlsx")
print("workbook written:", wb.sheetnames)